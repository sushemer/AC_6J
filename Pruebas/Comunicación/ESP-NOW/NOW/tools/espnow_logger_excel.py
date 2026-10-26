"""Registro serial ESP-NOW a JSONL y Excel, inspirado en gps_logger_excel.py."""
from __future__ import annotations

import argparse
import json
import queue
import re
import threading
import time
from datetime import datetime
from pathlib import Path

import serial
from serial.tools import list_ports
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ANSI = re.compile(r"\x1b\[[0-?]*[ -/]*[@-~]")
RX = re.compile(r"De ([0-9a-fA-F:]{17}) \| Seq: (\d+) \| RSSI: (-?\d+) dBm \| Msg: (.*)")
ESPNOW_RX = re.compile(
    r"ESPNOW_RX\s+role=(?P<role>\S+)\s+src=(?P<mac>[0-9a-fA-F:]{17})\s+"
    r"seq=(?P<seq>\d+)\s+rssi_dbm=(?P<rssi>-?\d+)\s+"
    r"rx_total=(?P<rx>\d+)\s+lost_total=(?P<lost>\d+)\s+"
    r"loss_pct=(?P<loss>[\d.]+)\s+queue_dropped=(?P<queue_dropped>\d+)\s+"
    r"callback_frames=(?P<callback_frames>\d+)\s+bad_len=(?P<bad_len>\d+)\s+"
    r"payload_bytes=(?P<payload_bytes>\d+)\s+esp_time_us=(?P<esp_time_us>-?\d+)\s+"
    r"rssi_min_dbm=(?P<rssi_min>-?\d+)\s+rssi_max_dbm=(?P<rssi_max>-?\d+)\s+"
    r"rssi_avg_dbm=(?P<rssi_avg>-?[\d.]+)\s+text=(?P<message>.*)"
)
STAMP = re.compile(r"\b[IWEDV] \((\d+)\)")
COLUMNS = ["fecha_pc", "rol", "puerto", "evento", "tiempo_sesion_s", "tiempo_esp_log_ms",
           "tiempo_esp_rx_us", "mac", "secuencia", "rssi_dbm", "rssi_min_dbm",
           "rssi_max_dbm", "rssi_avg_dbm", "paquetes_rx_total", "paquetes_perdidos_total",
           "perdida_pct", "descartes_cola", "callbacks_rx_total", "frames_longitud_invalida",
           "payload_bytes", "mensaje", "salto_secuencia",
           "estado_secuencia", "intervalo_rx_pc_ms", "procesamiento_python_ms",
           "linea_original", "bytes_hex"]


def parse_record(role, port, raw, elapsed, previous, received_ns=None):
    start = time.perf_counter_ns()
    line = ANSI.sub("", raw.decode("utf-8", errors="replace")).rstrip("\r\n")
    record = dict.fromkeys(COLUMNS)
    record.update(fecha_pc=datetime.now().astimezone().isoformat(timespec="milliseconds"),
                  rol=role, puerto=port, evento="log", tiempo_sesion_s=elapsed,
                  linea_original=line, bytes_hex=raw.hex())
    stamp = STAMP.search(line)
    if stamp:
        record["tiempo_esp_log_ms"] = int(stamp[1])
    match = ESPNOW_RX.search(line)
    if match:
        values = match.groupdict()
        mac = values["mac"]
        seq = int(values["seq"])
        key = (port, mac.lower())
        now_ns = received_ns if received_ns is not None else start
        last = previous.get(key)
        state = "primero"
        gap = 0
        if last:
            delta = (seq - last[0]) % (2**32)
            state = "duplicado" if delta == 0 else "retroceso_o_reinicio" if delta >= 2**31 else "avance"
            if state == "avance":
                gap = delta - 1
            record["intervalo_rx_pc_ms"] = (now_ns - last[1]) / 1_000_000
        previous[key] = (seq, now_ns)
        record.update(
            evento="recepcion", mac=mac, secuencia=seq,
            rssi_dbm=int(values["rssi"]), rssi_min_dbm=int(values["rssi_min"]),
            rssi_max_dbm=int(values["rssi_max"]), rssi_avg_dbm=float(values["rssi_avg"]),
            paquetes_rx_total=int(values["rx"]), paquetes_perdidos_total=int(values["lost"]),
            perdida_pct=float(values["loss"]), descartes_cola=int(values["queue_dropped"]),
            callbacks_rx_total=int(values["callback_frames"]),
            frames_longitud_invalida=int(values["bad_len"]),
            payload_bytes=int(values["payload_bytes"]), tiempo_esp_rx_us=int(values["esp_time_us"]),
            mensaje=values["message"], salto_secuencia=gap, estado_secuencia=state,
        )
    else:
        match = RX.search(line)
    if match and not ESPNOW_RX.search(line):
        mac, seq, rssi, message = match.groups()
        seq = int(seq)
        key = (port, mac.lower())
        now_ns = received_ns if received_ns is not None else start
        last = previous.get(key)
        state = "primero"
        gap = 0
        if last:
            delta = (seq - last[0]) % (2**32)
            state = "duplicado" if delta == 0 else "retroceso_o_reinicio" if delta >= 2**31 else "avance"
            if state == "avance":
                gap = delta - 1
            record["intervalo_rx_pc_ms"] = (now_ns - last[1]) / 1_000_000
        previous[key] = (seq, now_ns)
        record.update(evento="recepcion", mac=mac, secuencia=seq, rssi_dbm=int(rssi),
                      mensaje=message, salto_secuencia=gap, estado_secuencia=state)
    elif "Envio a " in line:
        record["evento"] = "resultado_envio"
        record["mensaje"] = line
    elif "Se perdieron " in line:
        record["evento"] = "alerta_perdida"
    elif "Error al invocar envio" in line:
        record["evento"] = "error_envio"
    record["procesamiento_python_ms"] = (time.perf_counter_ns() - start) / 1_000_000
    return record


def export_excel(journal, output):
    """Reconstruye el libro del registro durable; reemplaza solo al terminar."""
    book = Workbook(write_only=False)
    sheet = book.active
    sheet.title = "Registro"
    sheet.append(COLUMNS)
    with journal.open(encoding="utf-8") as source:
        for line in source:
            record = json.loads(line)
            sheet.append([record.get(column) for column in COLUMNS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    widths = [34, 14, 12, 22, 20, 22, 18, 22, 14, 14, 14, 14, 14, 18,
              20, 14, 16, 18, 22, 14, 34, 18, 26, 24, 28, 90, 50]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    # Mantener texto serial literal, incluso cuando empieza con '='.
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    temporary = output.with_suffix(".tmp.xlsx")
    try:
        book.save(temporary)
        temporary.replace(output)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)


def reader(role, port, events, stop):
    pending = bytearray()
    try:
        while not stop.is_set():
            chunk = port.read(port.in_waiting or 1)
            pending.extend(chunk)
            while b"\n" in pending:
                line, _, rest = pending.partition(b"\n")
                pending = bytearray(rest)
                events.put((role, port.port, bytes(line) + b"\n", time.perf_counter_ns(), None))
    except serial.SerialException as error:
        events.put((role, port.port, b"", time.perf_counter_ns(), str(error)))
    finally:
        if pending:
            events.put((role, port.port, bytes(pending), time.perf_counter_ns(), None))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--port", help="COM del receptor")
    parser.add_argument("--sender-port", help="COM del emisor, opcional")
    parser.add_argument("--baud", type=int, default=115200)
    parser.add_argument("--output", type=Path, help="Excel nuevo .xlsx")
    parser.add_argument("--save-seconds", type=float, default=30)
    parser.add_argument("--list-ports", action="store_true")
    parser.add_argument("--from-jsonl", type=Path, help="Exportar un registro previo sin hardware")
    args = parser.parse_args()
    if args.list_ports:
        for port in list_ports.comports():
            print(f"{port.device}: {port.description}")
        return
    if args.baud <= 0 or args.save_seconds <= 0:
        parser.error("Baud y save-seconds deben ser positivos")
    output = (args.output or Path("registros") / f"espnow_{datetime.now():%Y%m%d_%H%M%S}.xlsx").resolve()
    if output.suffix.lower() != ".xlsx":
        parser.error("La salida debe terminar en .xlsx")
    if output.exists():
        parser.error(f"La salida ya existe; elige otro nombre: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    if args.from_jsonl:
        export_excel(args.from_jsonl, output)
        print(f"Excel generado: {output}")
        return
    if not args.port:
        parser.error("Indica --port COM7 o usa --list-ports")
    if args.sender_port and args.sender_port.lower() == args.port.lower():
        parser.error("Emisor y receptor requieren puertos diferentes")
    journal = output.with_suffix(".jsonl")
    if journal.exists():
        parser.error(f"Ya existe el respaldo: {journal}")
    ports, threads = [], []
    events = queue.Queue()
    stop = threading.Event()
    previous = {}
    started = time.perf_counter_ns()
    last_save = time.monotonic()
    count = 0
    failed = False
    try:
        for role, name in [("receptor", args.port), ("emisor", args.sender_port)]:
            if name:
                port = serial.Serial(name, args.baud, timeout=0.2)
                ports.append((role, port))
        with journal.open("x", encoding="utf-8") as log:
            def store_event(event):
                nonlocal count, failed
                role, name, raw, received, error = event
                if error:
                    raw = f"ERROR SERIAL: {error}\n".encode()
                    failed = True
                    stop.set()
                record = parse_record(role, name, raw, (received - started) / 1e9, previous, received)
                log.write(json.dumps(record, ensure_ascii=False) + "\n")
                log.flush()
                count += 1
                print(f"[{role}] {record['linea_original']}")
            for role, port in ports:
                thread = threading.Thread(target=reader, args=(role, port, events, stop), daemon=True)
                thread.start()
                threads.append(thread)
            print(f"Registrando en {output}; Ctrl+C para terminar. Cierra los monitores seriales.")
            try:
                while not stop.is_set():
                    try:
                        store_event(events.get(timeout=0.2))
                    except queue.Empty:
                        pass
                    if time.monotonic() - last_save >= args.save_seconds:
                        try:
                            export_excel(journal, output)
                        except OSError as error:
                            print(f"No se pudo actualizar Excel: {error}. Los datos siguen en {journal}")
                        last_save = time.monotonic()
            except KeyboardInterrupt:
                print("Finalizando registro...")
            finally:
                stop.set()
                for thread in threads:
                    thread.join()
                while not events.empty():
                    store_event(events.get_nowait())
    except (serial.SerialException, OSError) as error:
        failed = True
        print(f"Error: {error}. Cierra cualquier monitor que use esos COM.")
    finally:
        stop.set()
        for _, port in ports:
            port.close()
        if journal.exists():
            try:
                export_excel(journal, output)
                print(f"Guardadas {count} líneas: {output}")
            except OSError as error:
                failed = True
                print(f"Excel pendiente: {error}. Recupera los datos con --from-jsonl {journal}")
    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
