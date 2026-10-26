#!/usr/bin/env python3
"""Guarda las métricas seriales emitidas por espnowClient y espnowSender.

Ejemplos:
    python espnow_metrics_logger.py --port COM7 --output mediciones.csv
    python espnow_metrics_logger.py --port COM7 --output mediciones.xlsx
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    import serial
except ImportError as exc:  # pragma: no cover - mensaje para uso en hardware
    raise SystemExit("Falta pyserial. Instale dependencias con: pip install -r requirements.txt") from exc


METRIC_RE = re.compile(
    r"ESPNOW_METRIC\s+role=(?P<role>\S+)\s+seq=(?P<seq>\d+)\s+"
    r"rssi_dbm=(?P<rssi>-?\d+)\s+rx=(?P<rx>\d+)\s+lost=(?P<lost>\d+)\s+"
    r"loss_pct=(?P<loss>[\d.]+)"
)
TX_RE = re.compile(
    r"ESPNOW_TX\s+seq=(?P<seq>\d+)\s+status=(?P<status>\S+)\s+"
    r"tx_ok=(?P<tx_ok>\d+)\s+tx_fail=(?P<tx_fail>\d+)"
)

FIELDS = (
    "timestamp_utc", "role", "seq", "rssi_dbm", "rx", "lost", "loss_pct",
    "status", "tx_ok", "tx_fail", "raw",
)


def parse_line(line: str) -> dict[str, str] | None:
    match = METRIC_RE.search(line)
    if match:
        values = match.groupdict()
        return {
            "role": values["role"], "seq": values["seq"],
            "rssi_dbm": values["rssi"], "rx": values["rx"],
            "lost": values["lost"], "loss_pct": values["loss"],
        }

    match = TX_RE.search(line)
    if match:
        values = match.groupdict()
        return {
            "role": "sender", "seq": values["seq"],
            "status": values["status"], "tx_ok": values["tx_ok"],
            "tx_fail": values["tx_fail"],
        }
    return None


def write_row(output: Path, row: dict[str, str]) -> None:
    row = {field: row.get(field, "") for field in FIELDS}
    row["timestamp_utc"] = datetime.now(timezone.utc).isoformat(timespec="milliseconds")
    row["raw"] = row.get("raw", "")

    if output.suffix.lower() == ".xlsx":
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError as exc:
            raise SystemExit("Para .xlsx instale openpyxl o use una salida .csv") from exc
        if output.exists():
            workbook = load_workbook(output)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(FIELDS)
        sheet.append([row[field] for field in FIELDS])
        workbook.save(output)
        return

    new_file = not output.exists() or output.stat().st_size == 0
    with output.open("a", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        if new_file:
            writer.writeheader()
        writer.writerow(row)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--port", required=True, help="Puerto serial, por ejemplo COM7")
    parser.add_argument("--baud", type=int, default=115200)
    parser.add_argument("--output", type=Path, default=Path("espnow_metrics.csv"))
    args = parser.parse_args()

    print(f"Escuchando {args.port} a {args.baud} baud; guardando en {args.output}")
    try:
        with serial.Serial(args.port, args.baud, timeout=1) as device:
            while True:
                raw = device.readline().decode("utf-8", errors="replace").strip()
                if not raw:
                    continue
                parsed = parse_line(raw)
                if parsed is None:
                    continue
                parsed["raw"] = raw
                write_row(args.output, parsed)
                print(f"guardado: {parsed['role']} seq={parsed['seq']}")
    except KeyboardInterrupt:
        print("\nLogger detenido.")
    except serial.SerialException as exc:
        print(f"Error serial: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
